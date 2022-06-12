import docx, xlsxwriter
import os, shutil, sys
from datetime import datetime
import re


maxWordsInFile = 300

class Paragraph:
    def __init__(self, nname, wwords) -> None:
        self.name = nname
        self.words = wwords
    pass

class Word:
    def __init__(self, enVal, ruVal, examp) -> None:
        self.enValue = enVal
        self.ruValue = ruVal
        self.example = examp
    pass
    
def getText(filename):
    doc = docx.Document(filename) 
    data = []
    wordsList = []

    usePar = False

    for para in doc.paragraphs:
        if len(para.text) == 0:
            continue
        cleanLine = re.sub('[^A-Za-z0-9]+', '', para.text)
        if len(cleanLine) == 0:
           continue
        if para.style.name == 'Heading 1':
            if usePar:
                wordsList.append(Word(enVal, ruVal, examp));
                data.append(Paragraph(parName, wordsList.copy()))
                wordsList.clear()
            
            enVal=ruVal=""    
            
            if para.text[0] != '#':
                parName = para.text
                usePar = True
            else:
                parName = ""
                usePar = False
            continue
        if para.style.name == 'Normal':
            if not usePar:
                continue
            if enVal and ruVal:
                wordsList.append(Word(enVal, ruVal, examp));
            if para.text == '\n':
                continue
            try:
                enVal, ruVal = para.text.split('\u2013')
            except Exception:
                print("Error splitting line:")
                print(para.text)
                exit()
            if enVal[-1] == ' ':
                enVal = enVal[:-1]
            if ruVal[0] == ' ':
                ruVal = ruVal[1:]
            ruVal = ruVal.capitalize()
            examp = ""
            continue
        if para.style.name == 'No Spacing':
            if not usePar:
                continue
            examp += para.text + '\n'
        else:
            continue

    if usePar:
         wordsList.append(Word(enVal, ruVal, examp));
         data.append(Paragraph(parName, wordsList.copy()))
         wordsList.clear()

    return data



# ru | en | part of speech | hint (example)
def writeText(data, outputDir, inputFileName, inputFile):

    current_datetime = datetime.now()
    current_datetime_text = "_" + str(current_datetime.year) +"_" + str(current_datetime.month) + "_" + str(current_datetime.day) + "_" + str(current_datetime.hour) + "_" + str(current_datetime.minute) + "_" + str(current_datetime.second)
    orig_name, end = inputFileName.split('.');

    outputDirName = outputDir + "Output_" + orig_name + current_datetime_text;

    os.mkdir(outputDirName)

    for paragraph in data:
        n = int(len(paragraph.words) / maxWordsInFile)
        for i in range(0, n + 1):
            outputFile = outputDirName + "/" + paragraph.name + "Part" + str(i + 1) + ".xlsx";
            workbook = xlsxwriter.Workbook(outputFile)
            worksheet = workbook.add_worksheet()

            row = 0
            for j in range(i * maxWordsInFile, (i + 1) * maxWordsInFile):
                if j == len(paragraph.words):
                    break
                worksheet.write(row, 0, paragraph.words[j].ruValue)
                worksheet.write(row, 1, paragraph.words[j].enValue)
                worksheet.write(row, 3, paragraph.words[j].example)
 
                row += 1

            workbook.close()
    shutil.copy(inputFile, outputDirName)
  
inputDir = sys.argv[1]
inputFileName = sys.argv[2]
outputDir = sys.argv[3]

inputFile = inputDir + inputFileName
data = getText(inputFile)

writeText(data, outputDir, inputFileName, inputFile)

print("Success!")